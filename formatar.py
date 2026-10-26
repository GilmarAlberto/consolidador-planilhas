from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = PatternFill("solid", fgColor="1F4E78")
CINZA = PatternFill("solid", fgColor="EEF2F7")
MOEDA = '"R$" #,##0.00'
DATA = "DD/MM/YYYY"
FONTE = "Calibri"
AZUL_CLARO = PatternFill("solid", fgColor="D6E4F0")
COLUNAS_SOMADAS = ("Quantidade", "Total")


def texto_exibido(celula, nome):
    """Como o valor aparece na tela, para calcular a largura certa."""
    valor = celula.value
    if valor is None:
        return ""
    if nome in ("Total", "Preço Unitário"):
        return f"R$ {valor:,.2f}"
    if nome == "Data":
        return "00/00/0000"
    return str(valor)


def formatar_relatorio(caminho):
    livro = load_workbook(caminho)

    for aba in livro.worksheets:
        # Cabeçalho: fundo azul, letra branca em negrito
        for celula in aba[1]:
            celula.fill = AZUL
            celula.font = Font(name=FONTE, bold=True, color="FFFFFF")
            celula.alignment = Alignment(horizontal="center")

        # Formato de cada coluna, pelo nome do cabeçalho
        for coluna in aba.iter_cols(min_row=1):
            nome = coluna[0].value
            for celula in coluna[1:]:
                celula.font = Font(name=FONTE)
                if nome in ("Total", "Preço Unitário"):
                    celula.number_format = MOEDA
                elif nome == "Data":
                    celula.number_format = DATA

            # Largura: o maior texto exibido + folga para a setinha do filtro
            maior = max(len(texto_exibido(c, nome)) for c in coluna[1:])
            maior = max(maior, len(nome) + 3)
            aba.column_dimensions[coluna[0].column_letter].width = maior + 6

        # Linhas zebradas para facilitar a leitura
        for linha in aba.iter_rows(min_row=3, max_row=aba.max_row):
            if linha[0].row % 2 == 1:
                for celula in linha:
                    celula.fill = CINZA

        aba.freeze_panes = "A2"               # cabeçalho fica fixo ao rolar
        aba.auto_filter.ref = aba.dimensions  # setinhas de filtro
        adicionar_linha_total(aba)

    livro.save(caminho)


def adicionar_linha_total(aba):
    """Linha TOTAL no fim da aba. Usa SUBTOTAL, que soma só as linhas
    visíveis: se o cliente filtrar uma loja, o total acompanha o filtro."""
    ultima = aba.max_row
    linha = ultima + 1
    borda = Border(top=Side(style="medium", color="1F4E78"))

    for coluna in range(1, aba.max_column + 1):
        celula = aba.cell(row=linha, column=coluna)
        nome = aba.cell(row=1, column=coluna).value
        letra = get_column_letter(coluna)
        if coluna == 1:
            celula.value = "TOTAL"
        elif nome in COLUNAS_SOMADAS:
            celula.value = f"=SUBTOTAL(109,{letra}2:{letra}{ultima})"
            if nome == "Total":
                celula.number_format = MOEDA
        celula.font = Font(name=FONTE, bold=True)
        celula.fill = AZUL_CLARO
        celula.border = borda
